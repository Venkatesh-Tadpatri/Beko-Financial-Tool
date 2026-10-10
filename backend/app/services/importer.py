"""Excel import: archive, read, map, validate.

Flow: upload (archive the file unchanged) -> choose sheet/header row -> map
columns -> stage every row -> validate -> user fixes or excludes rows ->
submit. Validation runs again at approval, against the master data of that
moment, before anything becomes reportable.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import stat
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core import fiscal
from app.core.datasets import AGEING_BUCKETS, DATASETS, DatasetDef, FieldDef
from app.core.money import MoneyParseError, to_count, to_paise
from app.models import Division, Product, SapCustomerMap, SapMaterialMap

MAX_FILE_BYTES = 100 * 1024 * 1024
MAX_ROWS = 500_000
PREVIEW_ROWS = 10


class ImportError_(ValueError):
    """User-facing problem with the file itself (not with individual rows)."""


# ------------------------------------------------------------------ files


def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _safe_name(name: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9._ -]+", "_", Path(name).name).strip() or "upload.xlsx"
    return stem[:120]


def archive_file(src: Path, uploads_dir: Path) -> tuple[str, str]:
    """Copy the original file, byte for byte, into the archive. Returns (relative path, sha256)."""
    if not src.is_file():
        raise ImportError_("The file could not be found.")
    if src.suffix.lower() != ".xlsx":
        raise ImportError_("Only Excel workbooks (.xlsx) can be imported.")
    size = src.stat().st_size
    if size > MAX_FILE_BYTES:
        raise ImportError_("The file is larger than 100 MB.")
    with open(src, "rb") as f:
        if f.read(2) != b"PK":
            raise ImportError_("The file is not a valid .xlsx workbook.")
    digest = sha256_of(src)
    rel = Path(digest[:2]) / f"{digest}_{_safe_name(src.name)}"
    target = uploads_dir / rel
    if not target.exists():
        target.parent.mkdir(parents=True, exist_ok=True)
        partial = target.with_name(target.name + ".partial")
        shutil.copyfile(src, partial)
        with open(partial, "rb+") as f:
            os.fsync(f.fileno())
        if sha256_of(partial) != digest:
            partial.unlink(missing_ok=True)
            raise ImportError_("The file changed while it was being copied. Please try again.")
        os.replace(partial, target)
        target.chmod(stat.S_IREAD)  # archived originals are read-only
    return rel.as_posix(), digest


# ------------------------------------------------------------------ reading


def _cell(v: object) -> object:
    """JSON-safe cell value that keeps numbers as numbers."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, float) and v != v:  # NaN
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    if isinstance(v, (int, float, bool)):
        return v
    return str(v)


def _headers(raw: tuple) -> list[str]:
    out: list[str] = []
    seen: dict[str, int] = {}
    for i, v in enumerate(raw, start=1):
        name = str(v).strip() if v not in (None, "") else f"Column {get_column_letter(i)}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        out.append(name)
    return out


def sheet_names(path: Path) -> list[str]:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for damaged files
        raise ImportError_("The workbook could not be opened. Is it a valid .xlsx file?") from exc
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def iter_sheet(path: Path, sheet: str, header_row: int) -> tuple[list[str], Iterator[tuple[int, dict]]]:
    """(headers, iterator of (excel_row_number, {header: value})). Empty rows are skipped."""
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        raise ImportError_(f'The sheet "{sheet}" does not exist in this workbook.')
    ws = wb[sheet]
    rows = ws.iter_rows(min_row=header_row, values_only=True)
    try:
        header_values = next(rows)
    except StopIteration:
        wb.close()
        raise ImportError_("The sheet is empty.") from None
    headers = _headers(header_values)

    def gen() -> Iterator[tuple[int, dict]]:
        try:
            for offset, values in enumerate(rows, start=1):
                cells = [_cell(v) for v in values[: len(headers)]]
                if all(c is None for c in cells):
                    continue
                yield header_row + offset, dict(zip(headers, cells, strict=False))
        finally:
            wb.close()

    return headers, gen()


def preview(path: Path, sheet: str, header_row: int) -> dict:
    headers, rows = iter_sheet(path, sheet, header_row)
    sample = []
    for _, r in rows:
        sample.append(r)
        if len(sample) >= PREVIEW_ROWS:
            break
    rows.close()  # type: ignore[attr-defined]
    return {"headers": headers, "sample": sample}


# ------------------------------------------------------------------ mapping suggestions

_SYNONYMS: dict[str, list[str]] = {
    "period_date": ["billing date", "posting date", "invoice date", "document date", "date", "fkdat", "budat", "period date"],
    "fiscal_year": ["fiscal year", "financial year", "fy", "year", "gjahr"],
    "period_no": ["posting period", "period", "month", "quarter", "monat", "poper"],
    "customer_code": ["sold to party", "sold-to party", "customer code", "customer", "customer number", "kunnr", "kunag", "payer"],
    "material_code": ["material", "material code", "material number", "matnr", "item code"],
    "amount": ["net value", "net amount", "amount", "sales value", "netwr", "value", "outstanding", "balance"],
    "quantity": ["billed quantity", "quantity", "qty", "fkimg"],
    "division_code": ["division", "division code", "spart"],
    "category": ["expense category", "category", "cost element", "gl account name"],
    "ageing_bucket": ["ageing bucket", "aging bucket", "ageing", "aging", "bucket", "days overdue"],
}


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def suggest_mapping(headers: list[str], ds: DatasetDef, template: dict[str, str] | None) -> dict[str, str]:
    if template and all(h in headers for h in template.values()):
        return dict(template)
    normalized = {_norm(h): h for h in headers}
    used: set[str] = set()
    result: dict[str, str] = {}
    for f in ds.fields:
        candidates = _SYNONYMS.get(f.key, []) + [_norm(f.label.split("(")[0]), f.key.replace("_", " ")]
        for cand in candidates:
            hit = normalized.get(_norm(cand))
            if hit and hit not in used:
                result[f.key] = hit
                used.add(hit)
                break
    return result


# ------------------------------------------------------------------ validation


@dataclass
class Issue:
    field: str | None
    code: str
    message: str
    severity: str = "error"  # error blocks submission; warning does not

    def as_dict(self) -> dict:
        return {"field": self.field, "code": self.code, "message": self.message, "severity": self.severity}


@dataclass
class Parsed:
    fy: int
    period_no: int
    dims: dict = field(default_factory=dict)  # resolved dimension values
    values: dict = field(default_factory=dict)  # field -> int (paise or count)
    refs: dict = field(default_factory=dict)  # resolved ids


@dataclass
class Lookups:
    customers: dict[str, int]  # lower(code) -> oem_id
    materials: dict[str, tuple[int, int, int]]  # lower(code) -> (product_id, segment_id, division_id)
    divisions: dict[str, int]  # lower(code) -> division_id

    @classmethod
    def load(cls, db: Session) -> "Lookups":
        customers = {c.lower(): o for c, o in db.execute(
            select(SapCustomerMap.customer_code, SapCustomerMap.oem_id).where(SapCustomerMap.is_active.is_(True)))}
        materials = {m.lower(): (p, s, d) for m, p, s, d in db.execute(
            select(SapMaterialMap.material_code, Product.id, Product.segment_id, Product.division_id)
            .join(Product, Product.id == SapMaterialMap.product_id).where(SapMaterialMap.is_active.is_(True)))}
        divisions = {c.lower(): i for c, i in db.execute(select(Division.code, Division.id).where(Division.is_active.is_(True)))}
        return cls(customers, materials, divisions)


_DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%dT%H:%M:%S", "%d-%b-%Y", "%d %b %Y")
_MONTHS = {m: i for i, m in enumerate(["apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec", "jan", "feb", "mar"], start=1)}


def parse_date(v: object) -> date:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 20000 <= v <= 80000:  # Excel serial date
            return date(1899, 12, 30) + timedelta(days=int(v))
        raise ValueError
    s = str(v).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s[:19] if "T" in s else s, fmt).date()
        except ValueError:
            continue
    raise ValueError


def parse_fiscal_year(v: object) -> int:
    if isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer():
        y = int(v)
    else:
        s = str(v).strip().upper().replace(" ", "")
        m = re.fullmatch(r"(?:FY)?(\d{4})(?:[-/](\d{2}|\d{4}))?", s)
        m2 = re.fullmatch(r"FY(\d{2})[-/](\d{2})", s)
        if m:
            y = int(m.group(1))
        elif m2:
            y = 2000 + int(m2.group(1))
        else:
            raise ValueError
    if not 1990 <= y <= 2100:
        raise ValueError
    return y


def parse_period_no(v: object, grain: str) -> int:
    if isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer():
        n = int(v)
    else:
        s = str(v).strip().lower()
        if grain == "M" and s[:3] in _MONTHS:
            n = _MONTHS[s[:3]]
        else:
            m = re.fullmatch(r"(?:q|p|h)?0*(\d{1,3})", s)
            if not m:
                raise ValueError
            n = int(m.group(1))
    if not 1 <= n <= fiscal.periods_per_year(grain):
        raise ValueError
    return n


def normalize_bucket(v: str) -> str:
    s = re.sub(r"\s+", "", v.lower()).replace("days", "")
    aliases = {"0-30": "0-30", "<30": "0-30", "<=30": "0-30", "1-30": "0-30", "31-60": "31-60", "61-90": "61-90",
               "90+": "90+", ">90": "90+", "91+": "90+", "91-": "90+", "over90": "90+", "morethan90": "90+"}
    if s in aliases:
        return aliases[s]
    raise ValueError


class Validator:
    def __init__(self, ds: DatasetDef, grain: str, column_map: dict[str, str], lookups: Lookups) -> None:
        self.ds, self.grain, self.column_map, self.lookups = ds, grain, column_map, lookups

    def value_of(self, f: FieldDef, raw: dict, edits: dict) -> object:
        if f.key in edits:
            v = edits[f.key]
            return None if v in (None, "") else v
        col = self.column_map.get(f.key)
        return raw.get(col) if col else None

    def _period(self, raw: dict, edits: dict) -> tuple[tuple[int, int] | None, list[Issue]]:
        d = self.value_of(self.ds.field("period_date"), raw, edits)
        fy_v = self.value_of(self.ds.field("fiscal_year"), raw, edits)
        pn_v = self.value_of(self.ds.field("period_no"), raw, edits)
        if d is not None:
            try:
                day = parse_date(d)
            except ValueError:
                return None, [Issue("period_date", "bad_date", f'"{d}" is not a date (use DD.MM.YYYY or YYYY-MM-DD).')]
            return (fiscal.fy_of(day), fiscal.period_of_month(self.grain, fiscal.fiscal_month_of(day))), []
        if fy_v is None:
            return None, [Issue("period_date", "missing_period", "Add a date, or a financial year and period.")]
        try:
            fy = parse_fiscal_year(fy_v)
        except ValueError:
            return None, [Issue("fiscal_year", "bad_fiscal_year", f'"{fy_v}" is not a financial year (e.g. 2025 or FY2025-26).')]
        if self.grain == "Y":
            return (fy, 1), []
        if pn_v is None:
            return None, [Issue("period_no", "missing_period", "Add the period number (month 1 = April, or quarter 1-4).")]
        try:
            return (fy, parse_period_no(pn_v, self.grain)), []
        except ValueError:
            limit = fiscal.periods_per_year(self.grain)
            return None, [Issue("period_no", "bad_period", f'"{pn_v}" is not a valid period (1-{limit}).')]

    def validate(self, raw: dict, edits: dict) -> tuple[Parsed | None, list[Issue]]:
        issues: list[Issue] = []
        period, p_issues = self._period(raw, edits)
        issues += p_issues
        dims: dict = {}
        values: dict = {}
        refs: dict = {}

        for f in self.ds.dimensions + self.ds.values:
            v = self.value_of(f, raw, edits)
            if v is None:
                if f.required:
                    issues.append(Issue(f.key, "missing", f"{f.label} is missing."))
                continue
            if f.kind == "text":
                dims[f.key] = str(v).strip() if not isinstance(v, float) or not v.is_integer() else str(int(v))
            else:
                try:
                    values[f.key] = to_paise(v) if f.kind == "money" else to_count(v)
                except MoneyParseError:
                    issues.append(Issue(f.key, "bad_number", f'{f.label}: "{v}" is not a number.'))
                    continue
                if values[f.key] < 0 and not f.negative_ok:
                    issues.append(Issue(f.key, "negative", f"{f.label} is negative. Check it is correct.", "warning"))

        if not any(f.key in values for f in self.ds.values):
            if not any(i.code in ("missing", "bad_number") for i in issues):
                issues.append(Issue(None, "no_values", "The row has no values."))

        # Master data lookups
        cust = dims.get("customer_code")
        if cust is not None:
            oem = self.lookups.customers.get(cust.lower())
            if oem is None:
                issues.append(Issue("customer_code", "unmapped_customer", f'SAP customer code "{cust}" is not mapped to an OEM.'))
            refs["oem_id"] = oem
        mat = dims.get("material_code")
        if mat is not None:
            hit = self.lookups.materials.get(mat.lower())
            if hit is None:
                issues.append(Issue("material_code", "unmapped_material", f'SAP material code "{mat}" is not mapped to a product.'))
            else:
                refs["product_id"], refs["segment_id"], refs["division_id"] = hit
        div = dims.get("division_code")
        if div is not None:
            did = self.lookups.divisions.get(div.lower())
            if did is None:
                issues.append(Issue("division_code", "unknown_division", f'Division "{div}" does not exist in master data.'))
            refs["division_id"] = did
        bucket = dims.get("ageing_bucket")
        if bucket is not None:
            try:
                dims["ageing_bucket"] = normalize_bucket(bucket)
            except ValueError:
                issues.append(Issue("ageing_bucket", "bad_bucket", f'"{bucket}" is not an ageing bucket ({", ".join(AGEING_BUCKETS)}).'))

        if period is None or any(i.severity == "error" for i in issues):
            return None, issues
        return Parsed(period[0], period[1], dims, values, refs), issues


# ------------------------------------------------------------------ whole-batch checks


def duplicate_issues(ds: DatasetDef, parsed: dict[int, Parsed]) -> dict[int, list[Issue]]:
    """Exact repeats -> warning (excluded by default). Same period+dimensions with different values -> error."""
    out: dict[int, list[Issue]] = {}
    exact: dict[tuple, int] = {}
    by_key: dict[tuple, int] = {}
    keyed = ds.fact_table != "sales"  # sales lines may legitimately repeat a customer+material in a month
    for row_no in sorted(parsed):
        p = parsed[row_no]
        dims = tuple(sorted(p.dims.items()))
        full = (p.fy, p.period_no, dims, tuple(sorted(p.values.items())))
        if full in exact:
            out.setdefault(row_no, []).append(Issue(None, "duplicate", f"Exact copy of row {exact[full]}.", "warning"))
            continue
        exact[full] = row_no
        if keyed:
            key = (p.fy, p.period_no, dims)
            if key in by_key:
                out.setdefault(row_no, []).append(Issue(
                    None, "conflict", f"Row {by_key[key]} already gives different values for the same period.", "error"))
            else:
                by_key[key] = row_no
    return out


def primary_value(ds: DatasetDef) -> str:
    return ds.values[0].key


def period_summary(ds: DatasetDef, grain: str, parsed: list[Parsed]) -> list[dict]:
    acc: dict[tuple[int, int], dict] = {}
    key = primary_value(ds)
    for p in parsed:
        a = acc.setdefault((p.fy, p.period_no), {"rows": 0, "total": 0})
        a["rows"] += 1
        a["total"] += p.values.get(key, 0)
    return [
        {"fy": fy, "period_no": n, "label": fiscal.period_label(fy, grain, n), "rows": a["rows"], "total_paise": a["total"]}
        for (fy, n), a in sorted(acc.items())
    ]


def months_covered(grain: str, periods: set[tuple[int, int]]) -> set[tuple[int, int]]:
    return {(fy, m) for fy, n in periods for m in fiscal.months_of(grain, n)}


def dataset(key: str) -> DatasetDef:
    if key not in DATASETS:
        raise ImportError_(f"Unknown dataset type {key!r}.")
    return DATASETS[key]


def dumps(v: object) -> str:
    return json.dumps(v, ensure_ascii=False, separators=(",", ":"))
